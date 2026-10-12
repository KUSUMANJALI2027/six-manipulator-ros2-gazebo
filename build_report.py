"""Build the project report and presentation from recorded, calculated results."""
from pathlib import Path
import json, csv
from html import escape
from docx import Document
from docx.shared import Inches, Pt
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle,Image,PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors
from reportlab.lib.utils import ImageReader
from pptx import Presentation
from pptx.util import Inches as PI, Pt as PP
ROOT=Path(__file__).resolve().parent;OUT=ROOT/'reports'
feedback=json.loads((OUT/'feedback_analysis.json').read_text());numbers=json.loads((OUT/'six_robot_numerical_results.json').read_text());valid=json.loads((OUT/'numerical_validation.json').read_text())
doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
styles=getSampleStyleSheet();styles['BodyText'].fontSize=9;styles['BodyText'].leading=13
story=[]
def heading(s,level=1):
    doc.add_heading(s,level);story.append(Paragraph(escape(s),styles['Heading'+str(min(level,3))]))
def para(s):
    doc.add_paragraph(s);story.append(Paragraph(escape(s),styles['BodyText']));story.append(Spacer(1,6))
def table(rows,widths=None):
    t=doc.add_table(rows=0,cols=len(rows[0]));t.style='Light Shading Accent 1'
    for row in rows:
        cells=t.add_row().cells
        for c,v in zip(cells,row):c.text=str(v)
    content=[[Paragraph(escape(str(v)),styles['BodyText']) for v in row] for row in rows]
    t=Table(content,colWidths=widths,repeatRows=1,hAlign='LEFT');t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e7eef5')),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),.25,colors.HexColor('#cbd5df')),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]));story.extend([t,Spacer(1,10)])
def page():doc.add_page_break();story.append(PageBreak())
def figure(path,width=480,maxheight=530):
    doc.add_picture(str(path),width=Inches(6.0));w,h=ImageReader(str(path)).getSize();scale=min(width/w,maxheight/h);story.append(Image(str(path),width=w*scale,height=h*scale));story.append(Spacer(1,8))
heading('Six-Manipulator Kinematics and Simulation Laboratory')
para('Project report | Vemula Kusumanjali | B.Tech ECE, CUTM Vizianagaram | 6 October 2026')
para('ROS 2 Jazzy and Gazebo Harmonic educational implementation: SCARA, Stanford-type, PUMA-type, SCORBOT-type, Delta and Stewart platform.')
heading('Abstract',2)
para('This project provides an interactive software laboratory for six serial and parallel robot geometries. It computes forward and inverse kinematics, displays joint or actuator variables and tool poses, samples position workspaces, and generates smooth joint-space waypoint trajectories. Four serial mechanisms use articulated Gazebo models with gravity and joint controllers. Delta and Stewart use a custom native-Bullet System plugin inside Gazebo, with rigid bodies and closed-loop constraints. Recorded feedback from all six robot families is included, with different test coverage stated explicitly. The implementation uses illustrative educational dimensions and supports simulation study without physical robot hardware.')
heading('Implementation status',2)
table([['Component','Evidence and current status'],['Six-robot FK/IK and trajectory generation','Implemented; reproducible numerical checks pass for all six models.'],['RViz and interactive control','User reported working motion and live values in the existing lab.'],['Four serial Gazebo models','Original feedback CSVs supplied for SCARA, Stanford, PUMA and SCORBOT.'],['Delta / Stewart articulated physics','Custom Bullet Gazebo integration verified for the prescribed small-motion trials; GUI screenshots received.'],['Collision-free motion and vendor fidelity','Not implemented; not claimed.']],[175,325])
para('The educational implementation now includes kinematics, workspace sampling, trajectories and physics-trial evidence for all six robot families. The results apply to the stated models and test conditions; they are not a general performance certification.')
page();heading('1. Objective and software architecture')
para('The laboratory answers three basic questions: where does the tool move for a specified set of joints; which actuator settings produce a requested tool pose; and which positions are achievable within the modeled limits? It also demonstrates smooth transitions between selected configurations and compares simulator feedback with commanded joint values.')
table([['Layer','Role'],['Python / NumPy / SciPy','Model geometry, transforms, inverse solvers and trajectory generation'],['ROS 2 Jazzy','MarkerArray, JointState and TF publication; simulator bridge'],['RViz','Robot geometry, workspace points and tool path visualization'],['Gazebo mirror','Externally computed bead geometry displayed using pose updates'],['Gazebo serial physics','Rigid bodies, gravity, PID joint control and feedback'],['Custom parallel physics','Native Bullet constraints stepped inside a Gazebo System plugin; simulated platform pose and actuator feedback'],['CSV analysis','Plots, command coverage, error metrics and source hashes']],[140,360])
para('The user environment is Ubuntu 24.04 on WSL2 with ROS 2 Jazzy and Gazebo Harmonic. The build environment used for this report executes numerical Python tests; it does not contain a running Gazebo or ROS installation. Simulator results therefore come from the original user-supplied CSV recordings.')
heading('2. Mathematical conventions')
para('All geometric lengths use metres. Internal joint rotations use radians. GUI target orientations use degrees. Serial robots use the standard DH transform A_i = Rz(theta_i) Tz(d_i) Tx(a_i) Rx(alpha_i). The complete tool transform is the ordered product A_1 ... A_n. The fourth column gives the tool position; the rotation block gives orientation.')
para('For a revolute joint the variable changes theta; for a prismatic joint it changes d. The geometry files specify the sign, offset and limits explicitly. Serial DH tables are supplied under dh_tables/. Parallel models use branch lengths and attachment geometry instead of an ordinary single-chain DH table.')
para('Serial IK uses bounded nonlinear least squares with multiple deterministic initial guesses. The residual contains position difference and a rotation-vector orientation difference. A failed solve means that no solution was found with the configured search and limits; it does not prove global unreachability.')
page();heading('3. Model definitions')
table([['Robot','Actuators','Model description'],['SCARA','R-R-P-R','0.30 / 0.25 m planar arms; downward vertical joint'],['Stanford','R-R-P-R-R-R','Educational telescoping arm with wrist'],['PUMA','6 revolute','Educational articulated arm with wrist'],['SCORBOT','5 revolute','Five-axis arm; arbitrary 6D poses may be infeasible'],['Delta','3 revolute inputs','Fixed-orientation platform; lower sphere-intersection branch'],['Stewart','6 leg lengths','Platform pose determined by six distance constraints']],[80,115,305])
para('For SCARA: x = 0.30 cos(q1) + 0.25 cos(q1+q2); y = 0.30 sin(q1) + 0.25 sin(q1+q2); z = 0.235 - q3; yaw = q1+q2+q4. Earlier height trials gave z=0.185 m at q3=0.050 m and z=0.135 m at q3=0.100 m.')
para('Delta: base radius 0.28 m, platform radius 0.075 m, upper arms 0.20 m, lower branches 0.48 m, and base elevation 0.85 m. Each branch constrains the distance between an elbow and a platform attachment. FK intersects three equal-radius spheres and chooses the lower assembly branch. The parallelogram model constrains orientation to zero.')
para('Stewart: base attachment radius 0.32 m, platform radius 0.20 m, home height 0.48 m, and leg strokes 0.30-0.75 m. Leg i has length ||p + R a_i - b_i||. IK evaluates these lengths directly. FK solves the six distance residuals locally near an upright seed; assembly selection and singularities remain relevant.')
para('All exact serial DH parameters, attachment angles, limits and reference positions are included in kinematics.py and model_parameters.json. The labels identify robot families; they are not claims of exact commercial robot dimensions.')
heading('4. Trajectories and workspace')
para('Each joint-space segment uses q(u)=q0+(q1-q0)(10u^3-15u^4+6u^5), where u ranges from 0 to 1. Velocity and acceleration are zero at segment endpoints. The minimum duration is increased to respect the configured per-actuator velocity and acceleration bounds. Every waypoint is a full stop. Joint-space interpolation does not generally produce a straight Cartesian line.')
para('Parallel loop closure is evaluated at 121 samples per planned segment. Workspace figures below are deterministic position samples, not collision-free volume estimates. Serial samples permit changing tool orientation. Stewart workspace samples use zero roll, pitch and yaw and therefore form a fixed-orientation subset.')
page();heading('5. Numerical verification')
para('The recovered model tests were rerun successfully. Each robot passed its home configuration and 12 near-home pose round trips, trajectory timing and limit checks, XML parsing of mirror geometry, and rejection of a distant test target. Four serial SDF models also passed coordinate mapping checks in 80 random configurations per robot and basic inertia/topology checks. These tests do not replace integrated simulator verification.')
para('The following values are calculated from the reproducible example trajectories in this release. IK residuals are mathematical consistency measures, not physical positioning accuracy.')
table([['Robot','Position residual (m)','Rotation residual (rad)','Workspace accepted / attempted']]+[[r['robot'],f"{r['ik_position_residual_m']:.3e}",f"{r['ik_orientation_residual_rad']:.3e}",f"{r['workspace_points']} / 1200"] for r in numbers],[75,130,130,165])
table([['Robot','Peak speed / limit','Peak acceleration / limit','Max sampled closure residual (m)']]+[[r['robot'],f"{r['max_velocity_limit_ratio']:.4f}",f"{r['max_acceleration_limit_ratio']:.4f}",'-' if r['max_sampled_loop_residual_m'] is None else f"{r['max_sampled_loop_residual_m']:.3e}"] for r in numbers],[75,130,130,165])
para('Ratios at or below one are consistent with the configured actuator limits for these calculated trajectories. The supplied trajectory CSVs explicitly label calculated positions, velocities and accelerations.')
for r in numbers:
    page();heading('Calculated result: '+r['robot']);figure(OUT/(r['robot']+'_workspace.png'),maxheight=390)
    para('Blue: modeled robot at home. Cyan: accepted position samples. Magenta: calculated tool trajectory. This figure is generated mathematically, not from a Gazebo dynamics recording.')
    para('Example target [x, y, z, roll, pitch, yaw] in metres/radians: '+', '.join(f'{v:.6f}' for v in r['target_pose']))
    para('Solved actuator coordinates in the model units: '+', '.join(f'{v:.6f}' for v in r['solved_actuators']))
page();heading('6. Recorded serial-robot physics results')
para('Gazebo models use illustrative masses and inertias, gravity, and joint position controllers. Command coordinates are converted to simulator joint offsets from the home configuration. Feedback is returned through the ROS-Gazebo bridge. The table summarizes the final recorded sample, not a guaranteed steady-state specification.')
table([['Robot','Samples','Time range (s)','Largest final revolute error (deg)','Final prismatic error (mm)']]+[[r['joints'][0]['robot'],r['samples'],f"{r['time_start_s']:.3f}-{r['time_end_s']:.3f}",f"{max(j['final_abs_error'] for j in r['joints'] if j['unit']=='deg'):.6f}",next((f"{j['final_abs_error']:.6f}" for j in r['joints'] if j['unit']=='mm'),'-')] for r in feedback],[65,55,110,140,130])
para('In all four selected recordings, only the J1 and J2 command columns varied; the other joints were held at constant commands. Holding-joint error is still useful evidence, but it is not dynamic excitation coverage of those axes. Prismatic residuals remain visible for SCARA and Stanford. No numerical acceptance tolerance was supplied by the instructor, so no global pass/fail accuracy threshold is invented.')
para('Error is measured minus command. sample_rmse is the square root of the mean squared error across recorded samples, including duplicate timestamps; it is not a time-weighted integral. Tool XYZ columns in these recordings were reconstructed using FK from measured joints and are not an independent tool-pose sensor measurement. CSV column alignment does not establish zero-latency command/feedback synchronization.')
for r in feedback:
    name=r['joints'][0]['robot'];page();heading('Recorded Gazebo trial: '+name);figure(OUT/(name+'_tracking.png'),maxheight=465)
    table([['Joint','Final |error|','Unit','Command changed?']]+[[j['joint'],f"{j['final_abs_error']:.6f}",j['unit'],'Yes' if j['command_varied'] else 'Held'] for j in r['joints']],[80,150,90,180])
    para('Source: '+r['source'])
parallel=json.loads((ROOT/'evidence'/'parallel'/'gazebo_verification.json').read_text())
page();heading('7. Verified parallel-robot Gazebo dynamics')
para('The user built libParallelBulletSystem.so and ran both worlds under Gazebo Sim 8.15.0. The uploaded run status is exit code 0, both plugin initialization messages are present, and no PARALLEL_ERROR messages were found. The CSV results were independently recalculated after upload. This is a native Bullet dynamics core integrated as a Gazebo System; it does not use the stock gz-physics adapter for these mechanisms.')
para('Delta uses three actuated upper-arm hinges and paired passive lower rods. Stewart uses six spherical-prismatic-spherical legs. Masses, cylinder-derived inertias, gravity and PID efforts determine motion; command positions are not imposed as body poses during integration. Bodies are advanced at internal steps of at most 1 ms, with contacts disabled. Initial five-second holds are followed by individual-actuator motions and return to home.')
table([['Robot','Recorded duration (s)','Samples','Actuators varied','Max sampled gap (mm)']]+[[r['robot'],f"{r['end_time_s']:.3f}",r['samples'],len(r['command_span']),f"{1000*r['max_anchor_or_slider_gap_m']:.6f}"] for r in parallel],[65,115,70,110,140])
table([['Robot','Max sampled pose/FK difference (mm)','Max orientation difference (rad)','Largest last-5s joint error']]+[[r['robot'],f"{1000*r['sampled_pose_vs_measured_actuator_fk_position_m']:.6f}",f"{r['sampled_pose_vs_measured_actuator_fk_rotation_rad']:.3e}",f"{max(r['last_5s_max_abs_error']):.3e} {r['units']}"] for r in parallel],[65,155,145,135])
para('Developer-selected investigation gates: sampled anchor/slider gap below 1 mm; sampled platform/FK position difference below 1 mm; orientation difference below 0.01 rad; final five-second tracking below 0.01 rad for Delta or 1 mm for Stewart. All gates passed, including actuator-limit and full-recording checks. These thresholds are screening criteria, not industrial accuracy specifications. Peak transient joint errors were approximately 0.01414 rad for Delta and 0.001336 m for Stewart, including startup gravity settling.')
para('Parallel platform poses are taken directly from the simulated platform rigid body. This differs from the earlier serial CSV tool coordinates, which were reconstructed from measured joints. Screenshots confirm rendering; motion and numerical claims are supported by the recordings. The separate local actuator-disabled gravity checks are retained in the add-on reference results.')
for r in parallel:
    page();heading(r['robot']+' — Gazebo screenshot and measured response')
    figure(ROOT/'evidence'/'screenshots'/(r['robot']+'_gazebo.png'),maxheight=265)
    para('User-provided Gazebo screenshot. Visual geometry is educational and does not include detailed fixtures or vendor CAD.')
    figure(OUT/(r['robot']+'_parallel_tracking.png'),maxheight=310)
    para('Recorded custom-physics response: command and measured actuator coordinates. Joint-error metrics are calculated from the full uploaded data; plotting only reduces visual density.')
page();heading('8. Limits of the evidence')
para('Delta and Stewart closed-loop dynamics ran successfully through the custom Bullet Gazebo System plugin. This bypasses the stock engine-adapter loop limitation without modifying the installed adapter. It is a custom dynamics integration, and the previous FK-only Gazebo mirror remains a separate visualization option. Robot contacts are disabled in the parallel trials.')
para('The four serial trials establish movement and feedback at the tested configurations. They do not establish all-joint dynamic coverage, stability for all trajectories, payload robustness, contact realism, actuator identification, or hardware accuracy. Joint and tool collision checks, singularity-aware planning and parallel passive-joint angle limits are not implemented.')
para('The parallel trials exercised all three Delta actuators and all six Stewart leg lengths, individually, around home. Finite data, measured limits, sampled anchor/slider gap, platform-pose versus FK agreement and final tracking passed the declared screening thresholds. Broader workspace, payload, disturbance and collision testing remains outside this evidence. Stewart uses an educational SPS leg realization, with no universal-joint angle limits.')
heading('9. Reproduction and demonstration')
para('Run bash start.sh in the isolated release folder. Choose 1 for kinematics, 2 for serial physics, or 4 for this report. For kinematics, select each robot, load demo waypoints, run waypoints and sample its workspace. The Gazebo mirror is optional geometry visualization. For the physics demonstration, select a serial robot, start Gazebo, wait for fresh feedback, and use the small test. Existing CSV evidence is already included, so repeats are not required merely to recreate the figures.')
para('Reproduce the numerical evidence with python3 validate_release.py, python3 analyze_feedback.py, and python3 export_six_robot_results.py. Report regeneration additionally uses python-docx, reportlab and python-pptx. Original recordings and SHA-256 hashes are retained for traceability.')
para('This release changes serial trajectory progress to simulator timestamps and isolates Gazebo transport sessions. The original CSVs document the earlier working runtime; the revised launch behavior has not been executed in the report-build environment. The parallel add-on compiled and ran on the user WSL machine, and Delta/Stewart GUI rendering is documented by screenshots. The original user project is not overwritten by the supplied installation instructions.')
heading('10. Software references')
para('Project primary sources: kinematics.py, serial_sdf.py, app.py, physics_lab.py, and the four original CSV recordings distributed with this release.')
para('ROS 2 Jazzy documentation: https://docs.ros.org/en/jazzy/ ; Gazebo Harmonic documentation: https://gazebosim.org/docs/harmonic/ ; SciPy least_squares and Rotation API documentation: https://docs.scipy.org/doc/scipy/reference/ . These document the software tools; project-specific numerical results come from the supplied implementation and recordings.')
doc.save(OUT/'Project_Report.docx')
def footer(canvas,document):
    canvas.setFont('Helvetica',8);canvas.setFillColor(colors.grey);canvas.drawString(40,22,'Six Robot Lab | Educational simulation | 6 October 2026');canvas.drawRightString(555,22,str(document.page))
SimpleDocTemplate(str(OUT/'Project_Report.pdf'),rightMargin=40,leftMargin=40,topMargin=35,bottomMargin=38).build(story,onFirstPage=footer,onLaterPages=footer)
# Concise editable presentation
prs=Presentation();prs.slide_width=PI(13.333);prs.slide_height=PI(7.5)
def slide(title,lines,image=None):
    s=prs.slides.add_slide(prs.slide_layouts[5]);s.shapes.title.text=title
    for p in s.shapes.title.text_frame.paragraphs:p.font.size=PP(28)
    box=s.shapes.add_textbox(PI(.55),PI(1.25),PI(5.5 if image else 12),PI(5.7));tf=box.text_frame;tf.word_wrap=True
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.size=PP(21);p.space_after=PP(17)
    if image:
        w,h=ImageReader(str(image)).getSize();scale=min(6/w,5.6/h);s.shapes.add_picture(str(image),PI(6.7),PI(1.35),width=PI(w*scale),height=PI(h*scale))
slide('Six-Manipulator Kinematics and Simulation Laboratory',['Vemula Kusumanjali | B.Tech ECE | CUTM Vizianagaram','ROS 2 Jazzy + Gazebo Harmonic','Educational geometry; no physical robot hardware','Progress and evidence review | 6 October 2026'])
slide('Project objective',['Joint inputs -> tool pose: forward kinematics','Target tool pose -> actuator values: inverse kinematics','Workspace: sampled reachable positions','Trajectory: smooth movement between selected waypoints'])
slide('Six robot families',['Serial: SCARA, Stanford, PUMA and SCORBOT','Parallel: Delta and Stewart platform','Serial models use standard DH transforms','Parallel models use attachment geometry and closure equations'])
slide('Software and outputs',['Python / NumPy / SciPy: numerical models','ROS 2: state, markers and TF','RViz: live geometry, workspace and path','Gazebo: serial physics + custom parallel Bullet dynamics'])
slide('Workspace and smooth trajectories',['Quintic interpolation with endpoint stops','Configured actuator speed / acceleration bounds','Joint-space motion can produce a curved tool path','Workspace samples do not include collision filtering'],OUT/'SCARA_workspace.png')
slide('Parallel-robot kinematics',['Delta: XYZ motion, fixed platform orientation','Stewart: six leg lengths and local forward closure','Calculated closure and trajectories verified','Custom Bullet Gazebo trials passed for both parallel models'],OUT/'Stewart_workspace.png')
slide('Numerical verification',['All six models passed numerical checks','Home + 12 near-home pose round trips per model','Four serial SDF mappings checked at 80 configurations each','Results are mathematical consistency evidence, not hardware accuracy'])
slide('Four recorded Gazebo physics trials',['SCARA, Stanford, PUMA and SCORBOT feedback analyzed','Recorded commands varied J1 and J2; other joints held','Per-joint errors and plots included in report','No all-joint or all-payload validation claim'],OUT/'SCARA_tracking.png')
slide('Current limitations',['Parallel validation covers prescribed small motions only','Residual prismatic offsets in SCARA and Stanford','Illustrative geometry and masses; no vendor fidelity','Collision-aware planning and singularity avoidance absent'])
slide('Demonstration sequence',['Show live FK, IK and DH/geometry values','Run a waypoint trajectory and workspace display','Show recorded command-versus-feedback plots','Show Delta and Stewart verified custom-physics demonstrations'])
slide('Parallel Gazebo trial evidence',['Delta: 60 s, all 3 inputs exercised','Stewart: 96 s, all 6 leg lengths exercised','Declared tracking, limits and closure screening gates passed','Custom native-Bullet System plugin; contacts disabled'],ROOT/'evidence'/'screenshots'/'Stewart_gazebo.png')
prs.save(OUT/'Project_Presentation.pptx')
print('Report PDF, editable DOCX and 11-slide presentation generated.')
